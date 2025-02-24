from docx import Document
import tkinter as tk
from tkinter import messagebox, ttk
import os
from Modules.utils import format_date, get_cell_value, clear_ui, select_excel_file

word_file = "Templates/Word_templates/Act_template.docx"

main_data = [
    ({'cell_row': '1'}, {'name': '_OBJECT-NAME_', 'cell_row': '1', 'cell_column': 'B'}),
    ({'cell_row': '2'}, {'name': '_CONTRACTOR-REPRESENTATIVE_', 'cell_row': '2', 'cell_column': 'B'}),
    ({'cell_row': '3'}, {'name': '_CONTRACTOR-REPRESENTATIVE-NAME_', 'cell_row': '3', 'cell_column': 'B'}),
    ({'cell_row': '4'}, {'name': '_TECHNICAL-SUPERVISION-REPRESENTATIVE_', 'cell_row': '4', 'cell_column': 'B'}),
    ({'cell_row': '5'}, {'name': '_TECHNICAL-SUPERVISION-REPRESENTATIVE-NAME_', 'cell_row': '5', 'cell_column': 'B'}),
    ({'cell_row': '6'}, {'name': '_DESIGN-ORGANIZATION-REPRESENTATIVE_', 'cell_row': '6', 'cell_column': 'B'}),
    ({'cell_row': '7'}, {'name': '_DESIGN-ORGANIZATION-REPRESENTATIVE-NAME_', 'cell_row': '7', 'cell_column': 'B'}),
    ({'cell_row': '8'}, {'name': '_ADDITIONAL-REPRESENTATIVES_', 'cell_row': '8', 'cell_column': 'B'}),
    ({'cell_row': '9'}, {'name': '_ADDITIONAL-REPRESENTATIVES-NAME_', 'cell_row': '9', 'cell_column': 'B'}),
    ({'cell_row': '10'}, {'name': '_GENERAL-CONTRACTOR_', 'cell_row': '10', 'cell_column': 'B'})
]

subobject_data = [
    {'name': '_SUBOBJECT-NAME_', 'cell_column': 'B', 'cell_row': '1'},
    {'name': '_ACT-NUMBER_', 'cell_column': 'A'},
    {'name': '_EXECUTION-DATE-MONTH_', 'cell_column': 'B'},
    {'name': '_EXECUTION-DATE_', 'cell_column': 'B'},
    {'name': '_WORK-NAMING_', 'cell_column': 'C'},
    {'name': '_ALBUM-NAME_', 'cell_column': 'D'},
    {'name': '_PAGE_', 'cell_column': 'E'},
    {'name': '_MATERIALS_', 'cell_column': 'F'},
    {'name': '_EXECUTIVE-DIAGRAM_', 'cell_column': 'G'},
    {'name': '_LABORATORY_', 'cell_column': 'H'},
    {'name': '_END-DATE_', 'cell_column': 'I'},
    {'name': '_NEXT-WORKS_', 'cell_column': 'J'}
]

class CreateActs:
    def __init__(self, root, menu_button, menu):
        self.root = root
        self.menu_button = menu_button
        self.menu = menu
        self.wb = None
        self.excel_file_name = None
        self.excel_file_path = None
        self.entries = {}
        self.sheet_vars = {}
        self.all_var = None

    def load_ui(self, project_name=None):
        wb, excel_file_name, excel_file_path = select_excel_file(project_name)

        if wb:
            self.wb = wb
            self.excel_file_name = excel_file_name
            self.excel_file_path = excel_file_path
            self.show_acts_ui()

    def show_acts_ui(self):
        clear_ui(self.root, self.menu_button, self.menu)
        main_sheet = self.wb['Main data']
        
        self.entries.clear()
        for i, (label_data, value_data) in enumerate(main_data, start=1):
            label_text = get_cell_value(main_sheet, label_data['cell_row'], 'A')
            value = get_cell_value(main_sheet, value_data['cell_row'], 'B')
            
            tk.Label(self.root, text=label_text).grid(row=i, column=0, padx=5, pady=5, sticky="e")
            self.entries[value_data['name']] = tk.Entry(self.root, width=100)
            self.entries[value_data['name']].grid(row=i, column=1, padx=5, pady=5)
            self.entries[value_data['name']].insert(0, value)

        frame = tk.Frame(self.root)
        frame.grid(row=11, column=0, columnspan=2, padx=5, pady=5, sticky="nsew")

        canvas = tk.Canvas(frame)
        scrollbar = ttk.Scrollbar(frame, orient="vertical", command=canvas.yview)
        scrollable_frame = tk.Frame(canvas)

        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )

        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        tk.Label(scrollable_frame, text="Выберите листы:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        self.sheet_vars.clear()
        for i, sheet_name in enumerate(self.wb.sheetnames, start=1):
            if sheet_name not in {'Main data', 'Contents'}:
                var = tk.BooleanVar()
                tk.Checkbutton(scrollable_frame, text=sheet_name, variable=var).grid(row=i, column=0, padx=5, pady=2, sticky="w")
                self.sheet_vars[sheet_name] = var

        canvas.grid(row=1, column=0, sticky="nsew")
        scrollbar.grid(row=1, column=1, sticky="ns")

        frame.grid_rowconfigure(1, weight=1)
        frame.grid_columnconfigure(0, weight=1)

        self.all_var = tk.BooleanVar()
        tk.Checkbutton(frame, text="Выбрать все", variable=self.all_var, command=self.toggle_all_sheets).grid(row=0, column=0, padx=5, pady=5, sticky="w")

        generate_button = tk.Button(self.root, text="Создать акты", command=self.generate_document)
        generate_button.grid(row=12, column=0, columnspan=2, pady=10)

    def save_to_excel(self):
        if self.wb is None:
            messagebox.showwarning("Предупреждение", "Сначала загрузите файл Excel!")
            return

        try:
            main_sheet = self.wb['Main data']
            for _, value_data in main_data:
                cell = f"{value_data['cell_column']}{value_data['cell_row']}"
                main_sheet[cell].value = self.entries[value_data['name']].get()
            
            self.wb.save(self.excel_file_path)
            messagebox.showinfo("Успех", f"Данные сохранены в {self.excel_file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить файл: {str(e)}")

    def generate_document(self):
        if self.wb is None:
            messagebox.showwarning("Предупреждение", "Сначала загрузите файл Excel!")
            return

        selected_sheets = [sheet_name for sheet_name, var in self.sheet_vars.items() if var.get()]
        if not selected_sheets:
            messagebox.showwarning("Предупреждение", "Выберите хотя бы один лист!")
            return

        self.save_to_excel()

        created_files = []
        errors = []

        base_dir = os.path.join(os.getcwd(), "Documents")
        os.makedirs(base_dir, exist_ok=True)

        excel_dir = os.path.join(base_dir, self.excel_file_name)
        os.makedirs(excel_dir, exist_ok=True)

        for sheet_name in selected_sheets:
            ws = self.wb[sheet_name]
            last_row = ws.max_row
            while last_row > 0 and all(cell.value is None for cell in ws[last_row]):
                last_row -= 1

            sheet_dir = os.path.join(excel_dir, sheet_name)
            os.makedirs(sheet_dir, exist_ok=True)

            for row_number in range(3, last_row + 1):
                replacements = {}
                for _, value_data in main_data:
                    replacements[value_data['name']] = self.entries[value_data['name']].get()

                for obj in subobject_data:
                    if obj['name'] == '_EXECUTION-DATE-MONTH_':
                        date_value = get_cell_value(ws, row_number, obj['cell_column'])
                        new_value = format_date(date_value) if date_value else ""
                        replacements[obj['name']] = new_value
                    elif obj['name'] == '_SUBOBJECT-NAME_':
                        replacements[obj['name']] = get_cell_value(ws, obj['cell_row'], obj['cell_column'])
                    else:
                        replacements[obj['name']] = get_cell_value(ws, row_number, obj['cell_column'])

                    if obj['name'] == '_ACT-NUMBER_':
                        act_number = get_cell_value(ws, row_number, obj['cell_column']).replace('/', '_')
                        output_path = os.path.join(sheet_dir, f"{act_number}.docx")

                try:
                    self.create_word_doc(word_file, output_path, replacements)
                    created_files.append(output_path)
                except Exception as e:
                    errors.append(f"Ошибка при создании {output_path}: {str(e)}")

        if created_files:
            success_message = "Успешно созданы акты:\n" + "\n".join(created_files)
            messagebox.showinfo("Успех", success_message)
        if errors:
            error_message = "Ошибки при создании актов:\n" + "\n".join(errors)
            messagebox.showerror("Ошибка", error_message)
        if not created_files and not errors:
            messagebox.showinfo("Информация", "Акт не были созданы.")

    def create_word_doc(self, template_path, output_path, replacements):
        doc = Document(template_path)
        for para in doc.paragraphs:
            for run in para.runs:
                for key, value in replacements.items():
                    if key in run.text:
                        run.text = run.text.replace(key, value)
        doc.save(output_path)

    def toggle_all_sheets(self):
        select_all = self.all_var.get()
        for var in self.sheet_vars.values():
            var.set(select_all)